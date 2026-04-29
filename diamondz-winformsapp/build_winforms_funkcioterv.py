from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = "Diamondz_WinForms_funkcioterv.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(9.5)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(header_cells[i], header, bold=True, color="FFFFFF")
        set_cell_shading(header_cells[i], "22294A")

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
            if len(table.rows) % 2 == 0:
                set_cell_shading(cells[i], "FAFAF8")

    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)

    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Segoe UI"
        run.font.color.rgb = RGBColor(34, 41, 74)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_note(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "E8E1D6")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Segoe UI"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(34, 41, 74)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(9.5)
    doc.add_paragraph()


def setup_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    styles["Normal"].font.name = "Segoe UI"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.08

    for style_name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        style = styles[style_name]
        style.font.name = "Segoe UI"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(34, 41, 74)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)


def build():
    doc = Document()
    setup_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(120)
    r = title.add_run("Diamondz")
    r.font.name = "Georgia"
    r.font.size = Pt(28)
    r.bold = True
    r.font.color.rgb = RGBColor(34, 41, 74)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Admin WinForms alkalmazás funkcióterv")
    r.font.name = "Segoe UI"
    r.font.size = Pt(18)
    r.bold = True
    r.font.color.rgb = RGBColor(202, 162, 107)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(18)
    r = meta.add_run("Készült a diamondz-winformsapp forráskódja alapján")
    r.font.name = "Segoe UI"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(90, 90, 90)

    doc.add_page_break()

    add_heading(doc, "1. Bevezetés")
    add_heading(doc, "1.1 Projekt célja", 2)
    doc.add_paragraph(
        "A projekt célja egy Windows Forms alapú adminisztrációs kliensalkalmazás fejlesztése a Diamondz webshop üzemeltetői számára. "
        "Az alkalmazás a meglévő DNN + Hotcakes rendszer REST végpontjaira épül, és áttekinthető felületet biztosít a termékek, készletadatok és kölcsönzések kezeléséhez."
    )
    doc.add_paragraph(
        "A WinForms kliens nem váltja ki a Hotcakes admin felületét, hanem célzott, gyors napi munkafelületet ad: dashboard mutatókat jelenít meg, listázza a termékeket és kölcsönzéseket, "
        "valamint lehetővé teszi a kiválasztott termék árának és készletének módosítását API-n keresztül."
    )

    add_heading(doc, "1.2 Érintett szereplők", 2)
    add_table(
        doc,
        ["Szereplő", "Leírás", "Hozzáférés"],
        [
            ["Adminisztrátor", "A webshop napi kezelője, aki termékeket, készletet és kölcsönzéseket ellenőriz.", "Diamondz Admin WinForms app"],
            ["Üzemeltető / fejlesztő", "API kulcsokat, végpontokat, telepítést és hibakeresést kezel.", "Forráskód, ApiSettings.cs, Visual Studio"],
            ["Rendszer", "DNN + Hotcakes webshop és a kölcsönzési modul API rétege.", "REST API, RentalApiHandler.ashx"],
        ],
        [4.0, 8.2, 5.2],
    )

    add_heading(doc, "1.3 Technológiai környezet", 2)
    add_table(
        doc,
        ["Terület", "Technológia / komponens"],
        [
            ["Kliensalkalmazás", ".NET 8 Windows Forms"],
            ["Fejlesztési nyelv", "C#"],
            ["Webshop motor", "DNN + Hotcakes"],
            ["Kommunikáció", "HTTP REST API, JSON válaszok"],
            ["Termék API", "Hotcakes /DesktopModules/Hotcakes/API/rest/v1/products"],
            ["Rendelés API", "Hotcakes orders végpont és Dnn.Kolcsonzes RentalApiHandler"],
            ["Adatmegjelenítés", "DataGridView, saját UserControl nézetek, dashboard kártyák és egyszerű diagramok"],
        ],
        [5.0, 11.5],
    )

    add_heading(doc, "2. Funkcionális leírás")
    add_heading(doc, "2.1 Főablak és navigáció", 2)
    doc.add_paragraph(
        "Az alkalmazás főablaka a MainForm.cs osztályban található. A felület bal oldali navigációs sávból és központi tartalompanelből áll. "
        "A menüpontok a Dashboard, Kölcsönzések és Termékek nézetek között váltanak, az aktív menüpont vizuálisan kiemelt állapotot kap."
    )
    add_bullets(
        doc,
        [
            "Maximalizált, reszponzívan méreteződő admin ablak.",
            "Bal oldali DIAMONDZ navigációs sáv.",
            "Dashboard, Kölcsönzések és Termékek nézetek elkülönített UserControl komponensekben.",
            "Globális Adatok újratöltése gomb, amely újra lekéri a termékeket és kölcsönzéseket.",
            "Betöltési állapot és hibaüzenet megjelenítése API hiba esetén.",
        ],
    )

    add_heading(doc, "2.2 Dashboard", 2)
    doc.add_paragraph(
        "A DashboardControl összesített mutatókat és gyors navigációs kártyákat jelenít meg. A kártyák kattinthatók, és közvetlenül a megfelelő termék- vagy kölcsönzéslistára viszik az adminisztrátort szűrt állapottal."
    )
    add_table(
        doc,
        ["Dashboard blokk", "Megjelenített adatok", "Művelet"],
        [
            ["Összes termék", "Termékek száma, elérhető termékek, teljes készlet darabszám", "Átkattintás a terméklistára"],
            ["Megvásárolható termékek", "Megvásárolható darabok, alacsony készlet, készlethiány", "Szűrt terméklista megnyitása"],
            ["Kölcsönözhető termékek", "Kölcsönzések száma, aktív, függőben lévő és befejezett kölcsönzések", "Szűrt kölcsönzéslista megnyitása"],
            ["Diagramok", "Fizetett / kosárban lévő kölcsönzések aránya, vásárolható készlet állapota", "Vizuális áttekintés"],
        ],
        [4.2, 8.0, 4.2],
    )

    add_heading(doc, "2.3 Termékek kezelése", 2)
    doc.add_paragraph(
        "A ProductsControl a Hotcakes termékkatalógus adatait jeleníti meg. A lista név és SKU alapján kereshető, gyorsszűrőkkel külön kezelhetők a bérelhető, megvásárolható és nem elérhető termékek."
    )
    add_bullets(
        doc,
        [
            "Terméknév, SKU, ár, készlet, elérhetőségi állapot és BVIN megjelenítése.",
            "Név és SKU szerinti szöveges szűrés.",
            "Gyorsszűrők: csak bérelhető, csak megvásárolható, nem elérhető termékek.",
            "Dashboardból érkező speciális szűrések: elérhető termékek, alacsony készlet, készlethiányos megvásárolható termékek.",
            "Rendezés oszlopfejlécre kattintással.",
            "CSV export a jelenlegi szűrt terméklistáról.",
            "Ár és raktárkészlet módosítása a kijelölt sorban, majd mentés API-n keresztül.",
        ],
    )
    add_note(
        doc,
        "Terméktípus felismerés",
        "A jelenlegi implementáció a termék nevében szereplő 'bérelhető' / 'berelheto' szöveg alapján különíti el a kölcsönözhető termékeket a megvásárolható termékektől.",
    )

    add_heading(doc, "2.4 Kölcsönzések kezelése", 2)
    doc.add_paragraph(
        "Az OrdersControl a kölcsönzési modul RentalApiHandler végpontjáról érkező foglalásokat jeleníti meg. A lista az adminisztrátori nyomon követést szolgálja: látható a vásárló, SKU, kezdő és záró dátum, API státusz és a számított aktuális kölcsönzési állapot."
    )
    add_table(
        doc,
        ["Funkció", "Leírás"],
        [
            ["Dátumszűrők", "Összes, aktív, ezen a héten, lejárt és függőben lévő kölcsönzések."],
            ["Rendezés", "Oszlopfejlécre kattintva növekvő/csökkenő rendezés."],
            ["Állapotkiemelés", "Aktív, függőben, befejezett és kosárban lévő státuszok színezése."],
            ["Részletező panel", "Vásárló, email, terméknév, SKU, napi díj, teljes díj, napok száma és státusz."],
            ["Másolás", "Terméknév és SKU gyors másolása admin munkafolyamatokhoz."],
        ],
        [4.5, 11.8],
    )

    add_heading(doc, "2.5 Készlet és kölcsönzési elérhetőség", 2)
    doc.add_paragraph(
        "Az alkalmazás betöltéskor lekéri a Hotcakes termékeket, a termékek inventory rekordjait és a kölcsönzési foglalásokat. "
        "A megvásárolható termékeknél a készletből levonja a már leadott, releváns rendelések mennyiségét. A kölcsönözhető termékeknél az aktív, fizetett foglalások alapján állítja be, hogy a termék elérhető-e."
    )
    add_bullets(
        doc,
        [
            "Aktív, Paid státuszú kölcsönzés esetén a bérelhető termék készlete 0-ra és nem elérhetőre áll.",
            "Nem aktív bérelhető termék esetén a kliens 1 darabos elérhetőséget számol.",
            "Megvásárolható termékeknél az elérhetőség a számított raktárkészlettől függ.",
            "Mentéskor a termék árát, státuszát, elérhetőségét és inventory mennyiségét frissíti a Hotcakes API-ban.",
        ],
    )

    add_heading(doc, "3. Adatmodell")
    add_heading(doc, "3.1 Product modell", 2)
    add_table(
        doc,
        ["Mező", "Típus", "Leírás"],
        [
            ["Bvin", "string", "Hotcakes termékazonosító."],
            ["Sku", "string", "Termék cikkszáma."],
            ["ProductName", "string", "Termék neve."],
            ["SitePrice", "decimal", "Webshopban használt ár."],
            ["Status", "int", "Hotcakes termék státusz."],
            ["IsAvailableForSale", "bool", "Hotcakes elérhetőségi jelző."],
            ["InventoryQuantity", "int", "A kliens által számított aktuális készlet."],
            ["IsRentableProduct", "bool", "Számított mező: név alapján bérelhető-e."],
            ["AvailabilityText", "string", "Elérhető / Nem elérhető megjelenítési szöveg."],
        ],
        [4.4, 3.0, 8.6],
    )

    add_heading(doc, "3.2 Order modell", 2)
    add_table(
        doc,
        ["Mező", "Típus", "Leírás"],
        [
            ["OrderNumber / Bvin / Id", "string / int", "Rendelés és foglalás azonosítói."],
            ["CustomerEmail / UserEmail", "string", "Vásárló e-mail címe, fallback logikával."],
            ["Sku", "string", "Kölcsönzött termék SKU-ja."],
            ["RentalStart / RentalEnd", "DateTime?", "Bérlés kezdete és vége a Hotcakes dátumformátumból parse-olva."],
            ["Status / StatusName", "string", "API-ból érkező státusz."],
            ["DailyPrice / TotalPrice", "decimal", "Napi díj és teljes bérleti díj."],
            ["CurrentRentalStatus", "string", "Számított állapot: Aktív, Függőben, Befejezett vagy kosárban lévő."],
            ["RentalDays", "int", "Bérlési napok száma."],
        ],
        [4.4, 3.2, 8.4],
    )

    add_heading(doc, "4. REST API kapcsolatok")
    add_table(
        doc,
        ["Végpont", "Metódus", "Feladat"],
        [
            ["/products?key=...", "GET", "Terméklista lekérése."],
            ["/products/{bvin}?key=...", "GET / POST", "Egy termék lekérése és frissítése."],
            ["/products/update?key=...", "POST", "Alternatív termékfrissítési végpont."],
            ["/orders?key=...", "GET", "Hotcakes rendelések lekérése a megvásárolható termékek készletszámításához."],
            ["/orders/{bvin}?key=...", "GET", "Rendelés részleteinek lekérése line item mennyiségekhez."],
            ["/productinventory...", "GET / POST", "Inventory rekordok lekérése és mentése."],
            ["/searchmanager/indexproduct...", "POST", "Termék újraindexelése mentés után."],
            ["/DesktopModules/MVC/Dnn.Kolcsonzes/RentalApiHandler.ashx", "GET", "Kölcsönzési foglalások lekérése API kulccsal."],
        ],
        [8.0, 2.3, 6.0],
    )
    add_note(
        doc,
        "API konfiguráció",
        "A BaseUrl és ApiKey értékek jelenleg az ApiSettings.cs fájlban vannak megadva. Éles környezetben javasolt ezeket konfigurációs fájlba vagy biztonságos titokkezelésbe áthelyezni.",
    )

    add_heading(doc, "5. Rendszerarchitektúra")
    add_heading(doc, "5.1 Komponensek", 2)
    add_table(
        doc,
        ["Komponens", "Leírás"],
        [
            ["Program.cs", "WinForms alkalmazás indítása és MainForm betöltése."],
            ["MainForm.cs", "Főablak, navigáció, adatbetöltés, mentési folyamat koordinálása."],
            ["DashboardControl.cs", "Összesített mutatók, kattintható kártyák és diagramok."],
            ["ProductsControl.cs", "Terméklista, szűrés, rendezés, CSV export, ár/készlet szerkesztés."],
            ["OrdersControl.cs", "Kölcsönzéslista, dátumszűrők, státuszszínezés, részletező panel."],
            ["HotcakesApiClient.cs", "HTTP kommunikáció, JSON feldolgozás, inventory és rendelési logika."],
            ["Models", "Product, Order, Address és ProductInventoryRecord adatmodellek."],
            ["HotcakesDateParser.cs", "Hotcakes /Date(...)/ formátumú dátumok kezelése."],
            ["ApiSettings.cs", "API végpontok és kulcsok központi összeállítása."],
        ],
        [5.0, 11.5],
    )

    add_heading(doc, "5.2 Adatfolyam", 2)
    add_numbered(
        doc,
        [
            "Az alkalmazás induláskor megjeleníti a Dashboard nézetet.",
            "A LoadDataAsync lekéri a termékeket a Hotcakes products végpontról.",
            "A kliens inventory adatokat keres termékenként vagy listás inventory API-ból.",
            "A kliens lekéri a kölcsönzéseket a RentalApiHandler.ashx végpontról.",
            "Az Order modellek kiszámítják az aktuális kölcsönzési állapotot az aktuális dátum alapján.",
            "A MainForm az aktív kölcsönzések alapján frissíti a bérelhető termékek elérhetőségét.",
            "A Dashboard, Products és Orders nézetek megkapják a frissített adatlistákat.",
            "Termék mentésekor a kliens frissíti a Hotcakes termékadatot, inventory rekordot, majd megpróbálja újraindexelni a terméket.",
        ],
    )

    add_heading(doc, "6. Hibakezelés és validáció")
    add_bullets(
        doc,
        [
            "API hiba esetén a felhasználó üzenetet kap, amely jelzi a BaseUrl és ApiKey ellenőrzésének szükségességét.",
            "Termékmentés előtt kötelező a BVIN megléte.",
            "Az árnak érvényes decimális számnak kell lennie.",
            "A készletnek egész számnak kell lennie, negatív érték nem menthető.",
            "A JSON válaszok feldolgozása rugalmas: a kliens több lehetséges gyökérkulcs alatt is megkeresi a listákat és objektumokat.",
            "Több alternatív Hotcakes végpont használata biztosítja, hogy különböző API konfigurációk mellett is működjön a mentés.",
        ],
    )

    add_heading(doc, "7. Telepítés és konfiguráció")
    add_heading(doc, "7.1 Fejlesztői indítás", 2)
    add_numbered(
        doc,
        [
            "A DiamondzWinForms.csproj fájl megnyitása Visual Studio-ban.",
            "A projekt futtatása Windows környezetben .NET 8 SDK-val.",
            "Az ApiSettings.cs fájlban a BaseUrl és ApiKey értékek ellenőrzése.",
            "Az alkalmazás indítása után az Adatok újratöltése gombbal ellenőrizhető az API kapcsolat.",
        ],
    )

    add_heading(doc, "7.2 Éles használati feltételek", 2)
    add_bullets(
        doc,
        [
            "Elérhető DNN + Hotcakes szerver és működő REST API.",
            "Érvényes Hotcakes API kulcs.",
            "Működő RentalApiHandler.ashx végpont a kölcsönzési adatokhoz.",
            "A szervernek engedélyeznie kell a kliens gépről érkező HTTP kéréseket.",
            "A terméknevekben következetesen szerepelnie kell a bérelhető jelölésnek, ha a jelenlegi felismerési logika marad.",
        ],
    )

    add_heading(doc, "8. Továbbfejlesztési lehetőségek")
    add_bullets(
        doc,
        [
            "API kulcs és BaseUrl áthelyezése külső konfigurációba.",
            "Külön terméktípus vagy custom property használata a bérelhető termékek felismeréséhez névalapú keresés helyett.",
            "Részletesebb kölcsönzés szerkesztési funkciók, például belső megjegyzés kezelése.",
            "Részletes naplózás API hibákhoz és sikertelen mentési kísérletekhez.",
            "Felhasználói jogosultságkezelés, ha több adminisztrátori szerepkör jelenik meg.",
            "Automatikus frissítés vagy időzített adatújratöltés dashboard használathoz.",
        ],
    )

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = footer.add_run("Diamondz Admin WinForms funkcióterv")
        r.font.name = "Segoe UI"
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(120, 120, 120)

    doc.save(OUT)


if __name__ == "__main__":
    build()
